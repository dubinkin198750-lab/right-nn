"""Разовый скрипт: переносит данные из старых Excel-таблиц (которые вели
вручную по каждому автору) в раздел «Блокировка» этого приложения —
создаёт дела блокировки (data/blocking_cases.json) по каждой ссылке.

ВАЖНО — что скрипт НЕ трогает:
- Листы «парсинг ...» в каждой книге (сырые результаты автоматического
  поиска по сайтам-складчинам) — не читаются и не импортируются вообще,
  это отдельный, самостоятельный процесс приложения (поиск/фильтры),
  сюда переносить нечего и не нужно.
- Раздел «Поиск» и всё, что с ним связано (data/authors.json,
  data/works.json, справочник сайтов) — скрипт только добавляет дела
  блокировки, никакие другие данные не изменяет.

Что переносится — по каждой строке каждого «рабочего» листа (не «парсинг»):
    Ссылка на произведение  -> url (обязательное поле, без него строка пропускается)
    Наличие в поисковой выдаче -> presence_google / presence_yandex
    Ответчик                -> defendant
    IP адрес                -> ip_address
    Подача мер / Дата подачи иска (в реальных таблицах там же зачастую
        записан номер и дата определения Мосгорсуда вперемешку, например
        «МГС 22.06Е 2И-5555 Удовл») / Принятое решение и дата / Комментарий /
        РКН / Повторная подача и дата / Исполнитель
                             -> одним текстом в notes (с подписями), чтобы
        ничего не потерять — эти поля в исходных таблицах не разделены на
        чистую дату/номер по отдельным ячейкам, а разложены человеком
        произвольно, автоматически раскладывать по датам-полям с типом
        <input type="date"> рискованно (испорченная дата в таком поле в
        интерфейсе просто не отобразится) — обычные текстовые "Примечания"
        безопаснее и не теряют ни одной детали, при этом сотрудник потом
        сам аккуратно разносит нужное по датам/номерам прямо в интерфейсе.

Работа автора (author_name) берётся не из «Автора» внутри таблицы (там он
проставлен только в отдельных строках, вперемешку с названиями
произведений — ненадёжно), а из явного соответствия файл → автор ниже,
теми же именами, что уже используются в остальном приложении (см.
backend/seed_new_authors.py) — чтобы дела сразу попали в ту же группу
автора в интерфейсе «Блокировка», а не создали случайно вторую параллельную
запись с чуть другим написанием имени.

Идемпотентность: storage.add_blocking_case() сама пропускает ссылку, если
она (по нормализованному виду) уже есть в блокировке — повторный запуск
скрипта безопасен, дублей не создаст.

Запуск из корня проекта, тем же python/venv, что и сам сервер:

    python -m backend.import_legacy_blocking_tables

По умолчанию ищет xlsx-файлы в папке legacy_tables/ рядом с проектом.
Свой путь можно указать явно:

    python -m backend.import_legacy_blocking_tables --dir /путь/к/таблицам
"""
import argparse
import os
import re

import openpyxl

from . import storage

_PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_DEFAULT_TABLES_DIR = os.path.join(_PROJECT_ROOT, "legacy_tables")

# Имя файла (как есть, без учёта регистра) -> каноническое имя автора,
# используемое во всём остальном приложении (см. seed_new_authors.py).
FILE_TO_AUTHOR = {
    "таблица_гофман.xlsx": "Гофман Ольга Сергеевна",
    "таблица_гордынец.xlsx": "Гордынец Иван Сергеевич",
    "таблица_белоусовой.xlsx": "Белоусова Алла Борисовна",
    "журнал_думай.xlsx": "НПМ Думай",
    "королева.xlsx": "Королева (Фелицына) Наталья",
    "таблица_федяев.xlsx": "Федяев Александр Алексеевич",
    "андрианов.xlsx": "Андрианов Евгений Владимирович",
}

# Возможные варианты подписи одной и той же колонки в разных файлах —
# таблицы велись годами разными людьми, заголовки чуть расходятся
# («Принятое решение и дата» / «Решение и дата» и т.п.). Проверяем
# вхождение подстроки, без учёта регистра.
_COL_URL = ["ссылка на произведение"]
_COL_PRESENCE = ["наличие в поисковой выдач"]  # см. _find_column — есть особый случай ниже
_COL_DEFENDANT = ["ответчик"]
_COL_IP = ["ip адрес", "ip-адрес"]
_COL_MEASURES = ["подача мер"]
_COL_CLAIM_DATE = ["дата подача иска", "дата подачи иска"]
_COL_DECISION = ["принятое решение и дата", "решение и дата"]
_COL_COMMENT = ["комментарий"]
_COL_RKN = ["ркн"]
_COL_REPEAT = ["повторная подача"]
_COL_EXECUTOR = ["исполнитель"]


def _norm(s):
    return (str(s) if s is not None else "").strip().lower()


def _find_column(header, needles, exclude_needles=None):
    """Первая колонка, чей заголовок содержит любую из needles (и не
    содержит ни одной из exclude_needles, если заданы) — без учёта
    регистра. Возвращает индекс или None.

    exclude_needles нужен для одного реального случая рассинхронизации
    заголовков в файле «Таблица_Гордынец.xlsx» (лист «Интеграция и обмен
    данными в 1С») — там заголовок колонки со ССЫЛКОЙ по ошибке содержит
    слова «Наличие в поисковой выдачи» внутри своего названия («Наличие в
    поисковой выдачи Ссылка на произведение...»), из-за чего простой
    поиск по подстроке для колонки «наличие» без исключения URL-словами
    находил бы не ту колонку."""
    for i, h in enumerate(header):
        hl = _norm(h)
        if not hl:
            continue
        if any(n in hl for n in needles) and not (exclude_needles and any(n in hl for n in exclude_needles)):
            return i
    return None


def _cell(row, idx):
    if idx is None or idx >= len(row):
        return ""
    value = row[idx]
    return "" if value is None else str(value).strip()


def _parse_presence(raw):
    low = raw.lower()
    return ("yandex" in low or "яндекс" in low), ("google" in low or "гугл" in low)


def _looks_like_url(value):
    return bool(re.match(r"^https?://", value.strip(), re.IGNORECASE))


def _build_notes(measures, claim_date_raw, decision, comment, rkn, repeat, executor):
    """Склеивает всё, что не укладывается в чистые поля формы, в одну
    читаемую заметку — с подписями, чтобы было понятно, откуда что взялось.
    Пустые куски пропускаются."""
    parts = []
    if measures:
        parts.append(f"Подача мер: {measures}")
    if claim_date_raw:
        parts.append(f"Дата/номер иска (МГС): {claim_date_raw}")
    if decision:
        parts.append(f"Решение и дата: {decision}")
    if comment:
        parts.append(f"Комментарий: {comment}")
    if rkn:
        parts.append(f"РКН: {rkn}")
    if repeat:
        parts.append(f"Повторная подача: {repeat}")
    if executor:
        parts.append(f"Исполнитель: {executor}")
    return " · ".join(parts)


def import_workbook(path, author_name, stats):
    wb = openpyxl.load_workbook(path, data_only=True)
    for sheet_name in wb.sheetnames:
        if "парсинг" in sheet_name.lower():
            continue  # раздел парсинга не трогаем — ни читаем, ни импортируем
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = list(rows[0])

        url_idx = _find_column(header, _COL_URL)
        if url_idx is None:
            print(f"    [пропускаю лист «{sheet_name}»] не нашёл колонку со ссылкой на произведение")
            continue
        presence_idx = _find_column(header, _COL_PRESENCE, exclude_needles=_COL_URL)
        defendant_idx = _find_column(header, _COL_DEFENDANT)
        ip_idx = _find_column(header, _COL_IP)
        measures_idx = _find_column(header, _COL_MEASURES)
        claim_date_idx = _find_column(header, _COL_CLAIM_DATE)
        decision_idx = _find_column(header, _COL_DECISION)
        comment_idx = _find_column(header, _COL_COMMENT)
        rkn_idx = _find_column(header, _COL_RKN)
        repeat_idx = _find_column(header, _COL_REPEAT)
        executor_idx = _find_column(header, _COL_EXECUTOR)

        for row in rows[1:]:
            url = _cell(row, url_idx)
            if not url or not _looks_like_url(url):
                continue

            presence_raw = _cell(row, presence_idx)
            presence_yandex, presence_google = _parse_presence(presence_raw)

            notes = _build_notes(
                _cell(row, measures_idx),
                _cell(row, claim_date_idx),
                _cell(row, decision_idx),
                _cell(row, comment_idx),
                _cell(row, rkn_idx),
                _cell(row, repeat_idx),
                _cell(row, executor_idx),
            )

            case = {
                "author_name": author_name,
                "work_title": sheet_name.strip(),
                "title": sheet_name.strip(),
                "source": f"импорт: {os.path.basename(path)} — «{sheet_name}»",
                "url": url,
                "presence_yandex": presence_yandex,
                "presence_google": presence_google,
                "defendant": _cell(row, defendant_idx),
                "ip_address": _cell(row, ip_idx),
                "notes": notes,
            }
            saved = storage.add_blocking_case(case)
            if saved is None:
                stats["duplicates"] += 1
            else:
                stats["imported"] += 1


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dir", default=_DEFAULT_TABLES_DIR, help="папка с xlsx-файлами (по умолчанию legacy_tables/ в корне проекта)")
    args = parser.parse_args()

    if not os.path.isdir(args.dir):
        print(f"Папка не найдена: {args.dir}")
        return

    stats = {"imported": 0, "duplicates": 0}
    for filename in sorted(os.listdir(args.dir)):
        if not filename.lower().endswith(".xlsx"):
            continue
        author_name = FILE_TO_AUTHOR.get(filename.lower())
        path = os.path.join(args.dir, filename)
        if not author_name:
            print(f"[пропускаю] «{filename}» — не знаю, какому автору соответствует "
                  f"(добавьте соответствие в FILE_TO_AUTHOR в этом скрипте)")
            continue
        print(f"Импортирую «{filename}» -> автор «{author_name}»...")
        before = dict(stats)
        import_workbook(path, author_name, stats)
        print(f"  добавлено: {stats['imported'] - before['imported']}, "
              f"уже было (пропущено как дубль): {stats['duplicates'] - before['duplicates']}")

    print(f"\nГотово. Всего добавлено дел: {stats['imported']}. "
          f"Пропущено как уже существующие: {stats['duplicates']}.")


if __name__ == "__main__":
    main()
