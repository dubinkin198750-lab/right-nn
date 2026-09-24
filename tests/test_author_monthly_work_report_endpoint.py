"""Тесты /api/authors/<id>/monthly-work-report.xlsx и
/api/settings/firm-letterhead."""
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tests._helpers import IsolatedStorageTestCase  # noqa: E402
from backend import storage  # noqa: E402


class TestAuthorMonthlyWorkReportEndpoint(IsolatedStorageTestCase):
    def setUp(self):
        super().setUp()
        from backend.app import app
        from backend import auth
        storage.create_user("admin1", auth.hash_password("x" * 10), "admin")
        storage.create_user("editor1", auth.hash_password("x" * 10), "editor")
        self.client = app.test_client()
        with self.client.session_transaction() as sess:
            sess["username"] = "admin1"
            sess["role"] = "admin"

    def test_downloads_valid_xlsx_for_existing_author(self):
        author = storage.upsert_author({"id": None, "name": "Автор"})
        resp = self.client.get(f"/api/authors/{author['id']}/monthly-work-report.xlsx")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data[:2], b"PK")
        self.assertIn("attachment", resp.headers["Content-Disposition"])

    def test_editor_cannot_download_report(self):
        """В акте печатаются коммерческие условия — только admin."""
        author = storage.upsert_author({"id": None, "name": "Автор"})
        with self.client.session_transaction() as sess:
            sess["username"] = "editor1"
            sess["role"] = "editor"
        resp = self.client.get(f"/api/authors/{author['id']}/monthly-work-report.xlsx")
        self.assertEqual(resp.status_code, 403)

    def test_404_for_unknown_author(self):
        resp = self.client.get("/api/authors/does-not-exist/monthly-work-report.xlsx")
        self.assertEqual(resp.status_code, 404)

    def test_rejects_bad_month_format(self):
        author = storage.upsert_author({"id": None, "name": "Автор"})
        resp = self.client.get(f"/api/authors/{author['id']}/monthly-work-report.xlsx?month=не-дата")
        self.assertEqual(resp.status_code, 400)

    def test_includes_works_and_their_cases(self):
        author = storage.upsert_author({"id": None, "name": "Автор"})
        work = storage.upsert_work({"id": None, "author_id": author["id"], "title": "Курс питона"})
        storage.add_blocking_case({
            "title": "т", "url": "https://pirate.test/x", "author_name": "Автор",
            "work_title": "Курс питона", "work_id": work["id"], "claim_date": "2026-08-05",
        })

        resp = self.client.get(f"/api/authors/{author['id']}/monthly-work-report.xlsx")
        self.assertEqual(resp.status_code, 200)

        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(resp.data))
        self.assertIn("Курс питона", wb.sheetnames)
        ws = wb["Курс питона"]
        found = any(
            cell == "https://pirate.test/x"
            for row in ws.iter_rows(values_only=True) for cell in row
        )
        self.assertTrue(found)

    def test_matches_cases_without_work_id_by_title_fallback(self):
        """Регрессия на найденный пробел: дела, добавленные через «Поиск
        по сайтам», могут иметь work_id=None — их всё равно нужно найти
        по автору+названию произведения."""
        author = storage.upsert_author({"id": None, "name": "Автор"})
        work = storage.upsert_work({"id": None, "author_id": author["id"], "title": "Курс питона"})
        storage.add_blocking_case({
            "title": "т", "url": "https://pirate.test/no-work-id", "author_name": "Автор",
            "work_title": "Курс питона", "work_id": None,  # как при добавлении через «Поиск по сайтам»
        })

        resp = self.client.get(f"/api/authors/{author['id']}/monthly-work-report.xlsx")
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(resp.data))
        ws = wb["Курс питона"]
        found = any(
            cell == "https://pirate.test/no-work-id"
            for row in ws.iter_rows(values_only=True) for cell in row
        )
        self.assertTrue(found)

    def test_uses_author_fields_in_act(self):
        """Реквизиты теперь в защищённой записи author_personal_data, не
        на самом объекте автора."""
        author = storage.upsert_author({"id": None, "name": "Автор"})
        from backend import author_personal_data
        author_personal_data.save(author["id"], {
            "customer_name": "ООО «Клиент»", "monthly_fee": "50000",
        })
        resp = self.client.get(f"/api/authors/{author['id']}/monthly-work-report.xlsx")
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(resp.data))
        ws = wb["Акт выполненных работ"]
        all_text = " ".join(str(cell) for row in ws.iter_rows(values_only=True) for cell in row if cell)
        self.assertIn("ООО «Клиент»", all_text)
        self.assertIn("50000", all_text)


class TestFirmLetterheadSettings(IsolatedStorageTestCase):
    def setUp(self):
        super().setUp()
        from backend.app import app
        from backend import auth
        storage.create_user("admin1", auth.hash_password("x" * 10), "admin")
        self.client = app.test_client()
        with self.client.session_transaction() as sess:
            sess["username"] = "admin1"
            sess["role"] = "admin"

    def test_default_is_empty(self):
        resp = self.client.get("/api/settings/firm-letterhead")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.get_json()["name"], "")

    def test_save_and_retrieve(self):
        resp = self.client.put("/api/settings/firm-letterhead", json={
            "name": "Юридический сервис Right-NN", "email": "antipiracy@right-nn.ru",
            "phone": "+7(831) 410-07-91", "address": "Нижний Новгород",
        })
        self.assertEqual(resp.status_code, 200)
        again = self.client.get("/api/settings/firm-letterhead").get_json()
        self.assertEqual(again["name"], "Юридический сервис Right-NN")
        self.assertEqual(again["email"], "antipiracy@right-nn.ru")

    def test_viewer_role_cannot_edit(self):
        with self.client.session_transaction() as sess:
            sess["role"] = "viewer"
        resp = self.client.put("/api/settings/firm-letterhead", json={"name": "X"})
        self.assertEqual(resp.status_code, 403)

    def test_viewer_role_cannot_view(self):
        # Реквизиты заказчика/юрфирмы — коммерчески чувствительная информация,
        # видна только администратору (см. заметки-к-развёртыванию.md).
        with self.client.session_transaction() as sess:
            sess["role"] = "viewer"
        resp = self.client.get("/api/settings/firm-letterhead")
        self.assertEqual(resp.status_code, 403)

    def test_editor_role_cannot_view(self):
        # До этого правки editor мог и читать, и менять реквизиты — теперь
        # только admin, даже на чтение.
        with self.client.session_transaction() as sess:
            sess["role"] = "editor"
        resp = self.client.get("/api/settings/firm-letterhead")
        self.assertEqual(resp.status_code, 403)


class TestAuthorNewFields(IsolatedStorageTestCase):
    def setUp(self):
        super().setUp()
        from backend.app import app
        from backend import auth
        storage.create_user("admin1", auth.hash_password("x" * 10), "admin")
        storage.create_user("editor1", auth.hash_password("x" * 10), "editor")
        self.client = app.test_client()
        with self.client.session_transaction() as sess:
            sess["username"] = "editor1"
            sess["role"] = "editor"

    def test_update_no_longer_accepts_customer_fields(self):
        """Осознанно изменённое поведение: реквизиты заказчика/договора
        больше не редактируются через обычное переименование автора (это
        мог делать любой editor) — перенесены в защищённую запись
        author_personal_data, доступную только admin."""
        author = storage.upsert_author({"id": None, "name": "Автор"})
        resp = self.client.put(f"/api/authors/{author['id']}", json={
            "name": "Автор", "customer_name": "ООО «Тест»", "monthly_fee": "15000",
        })
        self.assertEqual(resp.status_code, 200)
        data = resp.get_json()
        self.assertNotIn("customer_name", data)
        self.assertNotIn("monthly_fee", data)

    def test_customer_fields_now_saved_via_personal_data_admin_only(self):
        author = storage.upsert_author({"id": None, "name": "Автор"})
        resp = self.client.put(f"/api/authors/{author['id']}/personal-data", json={
            "customer_name": "ООО «Тест»", "monthly_fee": "15000",
        })
        self.assertEqual(resp.status_code, 403)  # editor, не admin

        with self.client.session_transaction() as sess:
            sess["username"] = "admin1"
            sess["role"] = "admin"
        resp2 = self.client.put(f"/api/authors/{author['id']}/personal-data", json={
            "customer_name": "ООО «Тест»", "monthly_fee": "15000",
        })
        self.assertEqual(resp2.status_code, 200)
        check = self.client.get(f"/api/authors/{author['id']}/personal-data").get_json()
        self.assertEqual(check["customer_name"], "ООО «Тест»")
        self.assertEqual(check["monthly_fee"], "15000")


class TestAuthorCustomerFieldsAdminOnly(IsolatedStorageTestCase):
    """Договорные условия с заказчиком — коммерчески чувствительная
    информация, живёт в той же защищённой записи, что и личные данные
    (author_personal_data) — видна и редактируема только роли admin."""
    def setUp(self):
        super().setUp()
        from backend.app import app
        from backend import auth
        storage.create_user("admin1", auth.hash_password("x" * 10), "admin")
        storage.create_user("editor1", auth.hash_password("x" * 10), "editor")
        self.client = app.test_client()
        self.author = storage.upsert_author({"id": None, "name": "Автор"})

    def _login_as(self, username, role):
        with self.client.session_transaction() as sess:
            sess["username"] = username
            sess["role"] = role

    def test_admin_can_set_customer_fields(self):
        self._login_as("admin1", "admin")
        resp = self.client.put(f"/api/authors/{self.author['id']}/personal-data", json={"customer_name": "ООО «Заказчик»"})
        self.assertEqual(resp.status_code, 200)
        check = self.client.get(f"/api/authors/{self.author['id']}/personal-data").get_json()
        self.assertEqual(check["customer_name"], "ООО «Заказчик»")

    def test_editor_cannot_set_customer_fields(self):
        self._login_as("editor1", "editor")
        resp = self.client.put(f"/api/authors/{self.author['id']}/personal-data", json={"customer_name": "ООО «Заказчик»"})
        self.assertEqual(resp.status_code, 403)

    def test_customer_fields_never_appear_on_plain_author_object(self):
        """Ключевая суть переноса: даже у admin эти поля больше не
        встречаются на самом объекте автора (GET /api/authors) — только
        через отдельный защищённый эндпоинт."""
        self._login_as("admin1", "admin")
        self.client.put(f"/api/authors/{self.author['id']}/personal-data", json={"customer_name": "ООО «Видно»"})
        resp = self.client.get("/api/authors")
        found = next(a for a in resp.get_json() if a["id"] == self.author["id"])
        self.assertNotIn("customer_name", found)


if __name__ == "__main__":
    unittest.main()
