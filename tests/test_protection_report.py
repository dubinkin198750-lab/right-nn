"""Отчёт для заказчика (23.09): ячейки должны совпадать с образцом
«Отчет по защите ИС в сети Интернет ООО Академия Маркетинга» буквально."""
import datetime as dt
import io
import os
import sys
import unittest

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tests._helpers import IsolatedStorageTestCase  # noqa: E402
from backend import reports, storage, author_personal_data  # noqa: E402


def ap(mgs_n, mgs_d, rkn_n, rkn_d, decision=""):
    return {"mgs_number": mgs_n, "mgs_date": mgs_d, "rkn_number": rkn_n, "rkn_date": rkn_d, "decision": decision}


# Строки 2, 17, 46, 82 образца
ROW2 = {"id": "r2", "url": "https://s138.skladchina.biz/threads/x.483624/", "claim_date": "2026-04-22",
        "appeals": [ap("2И-3828", "2026-04-24", "2026-04-29-458678", "2026-04-29", "заблокировано")],
        "block_date": "2026-05-15"}
ROW17 = {"id": "r17", "url": "https://slivup.info/threads/x/", "claim_date": "2026-04-22",
         "appeals": [ap("2И-3827", "2026-04-24", "2026-04-29-458675", "2026-04-29", "нет реакции"),
                     ap("2И-4753", "2026-05-26", "2026-05-28-471181", "2026-05-28", "заблокировано")],
         "block_date": "2026-06-04"}
ROW46 = {"id": "r46", "url": "https://big-money.net/threads/x/", "claim_date": "2026-04-28",
         "appeals": [ap("2И-3959", "2026-04-30", "2026-05-05-460633", "2026-05-05", "нет реакции"),
                     ap("2И-4756", "2026-05-26", "2026-05-28-471182", "2026-05-28", "нет реакции"),
                     ap("2И-5261", "2026-06-11", "2026-06-16-479496", "2026-06-16", "заблокировано")],
         "block_date": "2026-06-23"}
ROW82 = {"id": "r82", "url": "https://s60.zapret.me/threads/x/", "claim_date": "2026-06-09",
         "appeals": [ap("2И-5261", "2026-06-11", "2026-06-16-479496", "2026-06-16")]}


class TestRowsMatchSample(unittest.TestCase):
    def test_row2_first_appeal_blocked(self):
        r = reports.template_row(ROW2)
        self.assertEqual(r["B"], "Претензия администрации сайта 22.04")
        self.assertEqual(r["C"], "Определение Московского городского суда № 2И-3828 от 24.04.2026")
        self.assertEqual(r["D"], dt.date(2026, 5, 15))
        self.assertEqual((r["E"], r["F"]), ("", None))
        self.assertEqual(r["G"], "2026-04-29-458678 от 29.04")

    def test_row17_blocked_by_repeat_goes_to_second_date_column(self):
        r = reports.template_row(ROW17)
        self.assertIsNone(r["D"])
        self.assertEqual(r["E"], "Определение Московского городского суда № 2И-4753 от 26.05.2026")
        self.assertEqual(r["F"], dt.date(2026, 6, 4))
        self.assertEqual(r["G"], "2026-04-29-458675 от 29.04; 2026-05-28-471181 от 28.05")

    def test_row46_three_appeals(self):
        r = reports.template_row(ROW46)
        self.assertEqual(r["E"], "Определение Московского городского суда № 2И-4756 от 26.05.2026 / 2И-5261 от 11.06.2026")
        self.assertEqual(r["F"], dt.date(2026, 6, 23))
        self.assertEqual(r["G"], "2026-05-05-460633 от 05.05; 2026-05-28-471182 от 28.05; 2026-06-16-479496 от 16.06")

    def test_row82_not_blocked_no_dates(self):
        r = reports.template_row(ROW82)
        self.assertIsNone(r["D"])
        self.assertIsNone(r["F"])
        self.assertEqual(r["G"], "2026-06-16-479496 от 16.06")

    def test_period_filter(self):
        self.assertTrue(reports.case_in_period(ROW82, "2026-06-01", "2026-06-30"))
        self.assertFalse(reports.case_in_period(ROW82, "2026-07-01", "2026-07-31"))
        self.assertTrue(reports.case_in_period(ROW46, "2026-06-20", "2026-06-25"))  # только блокировка в периоде


def _load(xlsx_bytes):
    return openpyxl.load_workbook(io.BytesIO(xlsx_bytes))


class TestEndpoint(IsolatedStorageTestCase):
    def setUp(self):
        super().setUp()
        from backend.app import app
        from backend import auth
        storage.create_user("admin1", auth.hash_password("x" * 10), "admin")
        storage.create_user("editor1", auth.hash_password("x" * 10), "editor")
        self.client = app.test_client()
        with self.client.session_transaction() as s:
            s["username"], s["role"] = "admin1", "admin"
        self.a1 = storage.upsert_author({"id": None, "name": "Андрианов Евгений Владимирович"})
        self.a2 = storage.upsert_author({"id": None, "name": "Королева Наталья"})
        self.w1 = storage.upsert_work({"id": None, "author_id": self.a1["id"], "title": "Фабрика клонов", "query": "x",
                                       "customer_site_url": "https://academymarketing.ru/ai_klon"})
        self.w2 = storage.upsert_work({"id": None, "author_id": self.a1["id"], "title": "GPT's агенты", "query": "x"})
        self.w3 = storage.upsert_work({"id": None, "author_id": self.a2["id"], "title": "Прически для себя", "query": "x"})
        base = {"author_name": self.a1["name"], "work_title": "Фабрика клонов", "title": "t"}
        storage.add_blocking_case(dict(ROW82, **base))
        self.c17 = storage.add_blocking_case(dict(ROW17, **base, other_complaints=[
            {"id": "g1", "filed_at": "2026-05-08", "method": "Google DMCA", "notes": ""}]))
        # заблокированное дело — в архиве (как после автоархивации)
        storage.archive_reported_cases([dict(ROW2, **base)], 2026, 5)
        storage.add_blocking_case({"id": "k1", "url": "https://courses-free.ru/k/", "author_name": self.a2["name"],
                                   "work_title": "Прически для себя", "title": "t", "claim_date": "2026-06-10"})
        storage.add_blocking_case({"id": "g2", "url": "https://x.test/gpt", "author_name": self.a1["name"],
                                   "work_title": "GPT's агенты", "title": "t", "claim_date": "2026-09-01"})
        author_personal_data.save(self.a1["id"], {"customer_name": 'ООО "Академия маркетинга"',
                                                  "contract_number": "191/2024", "contract_date": "2026-04-23",
                                                  "monthly_fee": "39000"})

    def _get(self, **params):
        q = "&".join(f"{k}={v}" for k, v in params.items())
        r = self.client.get(f"/api/reports/protection.xlsx?{q}")
        if r.status_code != 200:
            self.fail(r.get_data(as_text=True)[:300])
        return _load(r.data)

    def test_author_report_sheets_named_like_sample(self):
        wb = self._get(date_from="2026-04-23", date_to="2026-07-23", author_id=self.a1["id"])
        self.assertEqual(wb.sheetnames, ['"Фабрика клонов" Яндекс', '"GPT\'s агенты" Яндекс',
                                         "Поисковая выдача Google", "Акт выполненных работ"])

    def test_archived_blocked_case_included_and_sorted(self):
        ws = self._get(date_from="2026-04-23", date_to="2026-07-23", author_id=self.a1["id"])['"Фабрика клонов" Яндекс']
        self.assertEqual([c.value for c in ws[1]], ["Ссылка", "Претензии", "Решения", "Дата блокировки",
                                                    "Повторное решение", "Дата блокировки", "Обращения Роскомнадзор"])
        urls = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
        self.assertEqual(urls, [ROW2["url"], ROW17["url"], ROW82["url"]])  # по дате претензии
        self.assertEqual(ws["D2"].value.date(), dt.date(2026, 5, 15))
        self.assertEqual(ws["D2"].number_format, "DD.MM.YYYY")
        self.assertEqual(ws["A2"].font.name, "Times New Roman")

    def test_period_excludes_other_months(self):
        ws = self._get(date_from="2026-09-01", date_to="2026-09-30", author_id=self.a1["id"])['"GPT\'s агенты" Яндекс']
        self.assertEqual(ws["A2"].value, "https://x.test/gpt")
        ws2 = self._get(date_from="2026-09-01", date_to="2026-09-30", author_id=self.a1["id"])['"Фабрика клонов" Яндекс']
        self.assertEqual(ws2.max_row, 1)  # только заголовок

    def test_single_work(self):
        wb = self._get(date_from="2026-04-01", date_to="2026-12-31", work_id=self.w3["id"])
        self.assertEqual(wb.sheetnames, ['"Прически для себя" Яндекс', "Поисковая выдача Google", "Акт выполненных работ"])
        self.assertEqual(wb.worksheets[0]["A2"].value, "https://courses-free.ru/k/")
        self.assertEqual(wb["Акт выполненных работ"]["C8"].value, "«Прически для себя»")

    def test_all_authors_together(self):
        wb = self._get(date_from="2026-04-01", date_to="2026-12-31")
        self.assertIn('"Прически для себя" Яндекс', wb.sheetnames)
        self.assertIn('"Фабрика клонов" Яндекс', wb.sheetnames)
        self.assertIn("Акт — Андрианов", wb.sheetnames)
        self.assertIn("Акт — Королева", wb.sheetnames)

    def test_google_sheet_like_sample(self):
        ws = self._get(date_from="2026-04-23", date_to="2026-07-23", author_id=self.a1["id"])["Поисковая выдача Google"]
        self.assertEqual(ws["A1"].value, ROW17["url"])
        self.assertEqual(ws["B1"].value, "Обращение от 08.05")

    def test_google_removed_date_marked_later(self):
        r = self.client.patch(f"/api/blocking-cases/{self.c17['id']}/other-complaints/g1", json={"resolved_at": "2026-05-12"})
        self.assertEqual(r.status_code, 200)
        ws = self._get(date_from="2026-04-23", date_to="2026-07-23", author_id=self.a1["id"])["Поисковая выдача Google"]
        self.assertEqual(ws["C1"].value.date(), dt.date(2026, 5, 12))

    def test_act_like_sample(self):
        ws = self._get(date_from="2026-04-23", date_to="2026-07-23", work_id=self.w1["id"])["Акт выполненных работ"]
        self.assertEqual(ws["C2"].value, "АКТ ВЫПОЛНЕННЫХ РАБОТ")
        self.assertEqual(ws["C4"].value, 'ООО "Академия маркетинга"')
        self.assertEqual(ws["C5"].value, "№ 191/2024 от «23» апреля 2026 года ")
        self.assertEqual(ws["C6"].value, "23.04.26 00:00 - 23.07.26 23:59")
        self.assertEqual(ws["C8"].value, "«Фабрика клонов»")
        self.assertIn("размещенное на сайте заказчика https://academymarketing.ru/ai_klon. Данный сайт", ws["B13"].value)
        self.assertEqual(ws["C27"].value, "Общая стоимость оказанных услуг составляет 39\u00a0000 рублей.")
        self.assertEqual(ws["C24"].value, reports.TEMPLATE_ACT_SERVICES[-1])

    def test_fee_from_author_data(self):
        ws = self._get(date_from="2026-04-23", date_to="2026-07-23", author_id=self.a1["id"])["Акт выполненных работ"]
        self.assertEqual(ws["C27"].value, "Общая стоимость оказанных услуг составляет 39\u00a0000 рублей.")

    def test_fee_corrected_in_dialog_goes_to_act_only(self):
        ws = self._get(date_from="2026-04-23", date_to="2026-05-23", author_id=self.a1["id"], fee="13000")["Акт выполненных работ"]
        self.assertEqual(ws["C27"].value, "Общая стоимость оказанных услуг составляет 13\u00a0000 рублей.")
        self.assertEqual(author_personal_data.get(self.a1["id"])["monthly_fee"], "39000")  # данные автора не тронуты

    def test_several_works_of_one_author(self):
        self.w4 = storage.upsert_work({"id": None, "author_id": self.a1["id"], "title": "Третье", "query": "x"})
        wb = self._get(date_from="2026-04-01", date_to="2026-12-31", work_ids=f"{self.w1['id']},{self.w2['id']}")
        self.assertEqual(wb.sheetnames, ['"Фабрика клонов" Яндекс', '"GPT\'s агенты" Яндекс',
                                         "Поисковая выдача Google", "Акт выполненных работ"])
        self.assertEqual(wb["Акт выполненных работ"]["C8"].value, "«Фабрика клонов», «GPT's агенты»")

    def test_works_of_different_authors(self):
        wb = self._get(date_from="2026-04-01", date_to="2026-12-31", work_ids=f"{self.w1['id']},{self.w3['id']}")
        self.assertIn('"Фабрика клонов" Яндекс', wb.sheetnames)
        self.assertIn('"Прически для себя" Яндекс', wb.sheetnames)
        self.assertNotIn('"GPT\'s агенты" Яндекс', wb.sheetnames)
        self.assertIn("Акт — Андрианов", wb.sheetnames)
        self.assertIn("Акт — Королева", wb.sheetnames)

    def test_unknown_work_in_list(self):
        r = self.client.get("/api/reports/protection.xlsx?date_from=2026-04-01&date_to=2026-12-31&work_ids=nope")
        self.assertEqual(r.status_code, 404)

    def test_links_without_results_not_in_report(self):
        storage.add_blocking_case({"url": "https://found-only.test/", "author_name": self.a1["name"],
                                   "work_title": "Фабрика клонов", "title": "t", "discovered_at": "2026-05-01"})
        ws = self._get(date_from="2026-04-23", date_to="2026-07-23", author_id=self.a1["id"])['"Фабрика клонов" Яндекс']
        urls = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
        self.assertNotIn("https://found-only.test/", urls)

    def test_block_dates_green_like_sample(self):
        ws = self._get(date_from="2026-04-23", date_to="2026-07-23", author_id=self.a1["id"])['"Фабрика клонов" Яндекс']
        self.assertEqual(ws["D2"].fill.fgColor.rgb, "FFADFF2F")
        self.assertIsNone(ws["D3"].value)
        self.assertNotEqual(ws["D3"].fill.fill_type, "solid")

    def test_act_logo_lines_letterhead(self):
        ws = self._get(date_from="2026-04-23", date_to="2026-07-23", work_id=self.w1["id"])["Акт выполненных работ"]
        self.assertEqual(len(ws._images), 1)
        self.assertEqual(ws["B12"].border.bottom.style, "thin")
        self.assertEqual(ws["B13"].border.top.style, "thin")
        self.assertIn("Right-NN", ws["C1"].value)  # реквизиты по умолчанию, если в настройках пусто
        self.assertIsNone(ws["B8"].alignment.horizontal)  # не по центру — не вылезает влево

    def test_title_already_in_quotes_not_doubled(self):
        w = storage.upsert_work({"id": None, "author_id": self.a2["id"], "query": "x",
                                 "title": "«Авторская система по созданию ИИ контент-завода «Фабрика клонов»"})
        wb = self._get(date_from="2026-04-01", date_to="2026-12-31", work_id=w["id"])
        self.assertEqual(wb["Акт выполненных работ"]["C8"].value, "«Авторская система по созданию ИИ контент-завода «Фабрика клонов»")
        self.assertTrue(wb.sheetnames[0].startswith('"Авторская система'))

    def test_bad_period_rejected(self):
        self.assertEqual(self.client.get("/api/reports/protection.xlsx?date_from=2026-05-01").status_code, 400)
        self.assertEqual(self.client.get("/api/reports/protection.xlsx?date_from=2026-06-01&date_to=2026-05-01").status_code, 400)

    def test_editor_cannot_download(self):
        with self.client.session_transaction() as s:
            s["username"], s["role"] = "editor1", "editor"
        r = self.client.get("/api/reports/protection.xlsx?date_from=2026-05-01&date_to=2026-05-31")
        self.assertEqual(r.status_code, 403)


if __name__ == "__main__":
    unittest.main()
